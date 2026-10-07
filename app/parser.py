"""简历文本提取：支持 PDF / DOCX / TXT / MD。

扫描件 / 图片型 PDF 无文本层，自动回退到本地 OCR（RapidOCR，离线运行）。
"""

from __future__ import annotations

import io

from fastapi import HTTPException

# OCR 最多处理页数与渲染倍率（2x ≈ 144dpi，兼顾速度与识别率）
OCR_PAGE_LIMIT = 12
OCR_ZOOM = 2

# 最小有效文本长度（字符），低于此值视为提取失败
MIN_TEXT_LEN = 50


class ParseError(HTTPException):
    def __init__(self, detail: str):
        super().__init__(status_code=422, detail=detail)


def _decode_text(data: bytes) -> str:
    """按 utf-8 → gb18030 → replace 的顺序解码，兼容中文 txt。"""
    for encoding in ("utf-8", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _extract_pdf_text_layer(data: bytes) -> str:
    """读取 PDF 自带文本层（普通 PDF）。扫描件会返回空字符串。"""
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [(page.extract_text() or "").strip() for page in reader.pages]
    except Exception:  # noqa: BLE001
        return ""
    return "\n".join(p for p in pages if p)


_OCR_ENGINE = None


def _get_ocr_engine():
    """懒加载并复用 OCR 引擎（首次加载模型约 1 秒）。"""
    global _OCR_ENGINE
    if _OCR_ENGINE is None:
        from rapidocr_onnxruntime import RapidOCR

        _OCR_ENGINE = RapidOCR()
    return _OCR_ENGINE


def _ocr_pdf(data: bytes) -> str:
    """将 PDF 逐页渲染为图片并 OCR 识别（阻塞式，调用方应放入线程）。"""
    import pymupdf  # PyMuPDF
    import numpy as np

    engine = _get_ocr_engine()
    texts: list[str] = []
    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"无法打开 PDF 文件：{exc}") from exc

    with doc:
        for i in range(min(doc.page_count, OCR_PAGE_LIMIT)):
            pix = doc[i].get_pixmap(matrix=pymupdf.Matrix(OCR_ZOOM, OCR_ZOOM), alpha=False)
            img = np.frombuffer(pix.samples, dtype=np.uint8).reshape(
                pix.height, pix.width, pix.n
            )
            # RapidOCR 基于 OpenCV，按 BGR 处理；PyMuPDF 输出 RGB，需翻转通道
            result, _ = engine(img[:, :, ::-1])
            if result:
                texts.append("\n".join(str(line[1]) for line in result))

    return "\n".join(t for t in texts if t.strip())


def _extract_docx(data: bytes) -> str:
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"DOCX 解析失败：{exc}") from exc

    parts: list[str] = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:  # 表格常见于简历（基本信息、技能清单）
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_text(filename: str, data: bytes) -> tuple[str, dict]:
    """根据扩展名分发到对应解析器。

    返回 (文本, 元数据)；元数据中的 ocr_used 表示是否使用了 OCR 回退。
    """
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    meta = {"ocr_used": False}

    if ext == ".pdf":
        text = _extract_pdf_text_layer(data)
        if len(text.strip()) < MIN_TEXT_LEN:
            # 无文本层或文本极少 → 本地 OCR 回退
            ocr_text = _ocr_pdf(data)
            if len(ocr_text.strip()) >= MIN_TEXT_LEN:
                text = ocr_text
                meta["ocr_used"] = True
    elif ext == ".docx":
        text = _extract_docx(data)
    elif ext in (".txt", ".md"):
        text = _decode_text(data)
    else:
        raise ParseError(f"不支持的文件格式「{ext or filename}」，请上传 PDF / DOCX / TXT / MD")

    text = text.strip()
    if len(text) < MIN_TEXT_LEN:
        raise ParseError(
            "未能从文件中提取到足够文本（扫描件 OCR 识别结果过少或内容为空），"
            "请确认文件内容，或改用 DOCX / TXT 格式"
        )
    return text, meta
